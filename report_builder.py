from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
D=Document(); sec=D.sections[0]; sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78);sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
NAVY='183A35';GREEN='315C4D';SAGE='EAF0E8';PALE='F5F7F3';GOLD='B98955';INK='26332E';MUTED='64736B';WHITE='FFFFFF'
styles=D.styles
for n,z,c in [('Normal',9.5,INK),('Title',29,NAVY),('Heading 1',17,NAVY),('Heading 2',12,GREEN),('Heading 3',10,GREEN)]:
 s=styles[n];s.font.name='Arial';s.font.size=Pt(z);s.font.color.rgb=RGBColor.from_string(c);s.font.bold=n!='Normal';s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=n.startswith('Heading')
for n in ['List Bullet','List Bullet 2','List Number']:styles[n].font.name='Arial';styles[n].font.size=Pt(9.5)
h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=h.add_run('RENTEASE  /  PROJECT REPORT');r.font.name='Arial';r.font.size=Pt(8);r.bold=True;r.font.color.rgb=RGBColor.from_string(GREEN)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=f.add_run('Azir Sam C  |  Web Development Internship  |  Page ');r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(MUTED);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
D.core_properties.title='RentEase Furniture and Appliance Rental Platform Project Report';D.core_properties.author='Azir Sam C';D.core_properties.subject='Detailed technical and functional project report'
def para(t='',style=None,center=False):
 p=D.add_paragraph(style=style) if style else D.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.1
 if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run(t);return p
def H(t,l=1):
 p=D.add_heading(t,level=l);p.paragraph_format.space_before=Pt(10 if l==1 else 7);p.paragraph_format.space_after=Pt(4);return p
def bullet(t):
 p=D.add_paragraph(t,style='List Bullet');p.paragraph_format.space_after=Pt(3);return p
def number(t):
 p=D.add_paragraph(t,style='List Number');p.paragraph_format.space_after=Pt(3);return p
def table(head,rows,widths=None,fs=8.2):
 t=D.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,x in enumerate(head):
  c=t.rows[0].cells[i];c.text=x;c._tc.get_or_add_tcPr().append(OxmlElement('w:shd'));c._tc.tcPr[-1].set(qn('w:fill'),NAVY)
  for r in c.paragraphs[0].runs:r.bold=True;r.font.name='Arial';r.font.size=Pt(fs);r.font.color.rgb=RGBColor.from_string(WHITE)
 for rownum,row in enumerate(rows):
  cells=t.add_row().cells
  for i,x in enumerate(row):
   c=cells[i];c.text=str(x);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if rownum%2:c._tc.get_or_add_tcPr().append(OxmlElement('w:shd'));c._tc.tcPr[-1].set(qn('w:fill'),PALE)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1)
    for r in p.runs:r.font.name='Arial';r.font.size=Pt(fs);r.font.color.rgb=RGBColor.from_string(INK)
  if widths:
   for i,w in enumerate(widths):cells[i].width=Inches(w)
 if widths:
  for i,w in enumerate(widths):
   for c in t.rows[0].cells:c.width=Inches(w)
 D.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
def callout(title,text):
 t=D.add_table(rows=1,cols=1);c=t.cell(0,0);c.text='';c._tc.get_or_add_tcPr().append(OxmlElement('w:shd'));c._tc.tcPr[-1].set(qn('w:fill'),SAGE)
 p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);r=p.add_run(title);r.bold=True;r.font.color.rgb=RGBColor.from_string(GREEN);r.font.size=Pt(10)
 p=c.add_paragraph(text);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
 for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string(INK)
 D.add_paragraph().paragraph_format.space_after=Pt(0)
# Cover
p=para('RENTEASE',center=True);p.paragraph_format.space_before=Pt(80);p.runs[0].font.size=Pt(13);p.runs[0].bold=True;p.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
p=D.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10);p.add_run('Furniture and Appliance\nRental Platform')
p=para('Detailed Project Report',center=True);p.runs[0].font.size=Pt(15);p.runs[0].font.color.rgb=RGBColor.from_string(GOLD)
t=D.add_table(rows=1,cols=1);c=t.cell(0,0);c.text='A responsive, multi-shop rental platform for urban living';c._tc.get_or_add_tcPr().append(OxmlElement('w:shd'));c._tc.tcPr[-1].set(qn('w:fill'),NAVY);c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string(WHITE);r.bold=True;r.font.size=Pt(11)
D.add_paragraph()
table(['Prepared by','Details'],[('Name','Azir Sam C'),('Department','Computer Science and Engineering (CSE)'),('Internship group','Web Development'),('Report date','26 October 2026')],[1.65,4.35],9)
para('PROJECT REPORT',center=True).runs[0].font.size=Pt(8)
D.add_page_break()
H('Report Guide')
para('This report describes the project purpose, implemented workflows, architecture, data design, security controls, deployment, and current limitations. It is prepared for academic review and technical handover.')
table(['Section','Coverage'],[('1. Executive Summary','Purpose, product scope, and technology overview'),('2. Background and Problem','Rental context and user pain points'),('3. Objectives and Scope','Goals, boundaries, assumptions, and constraints'),('4. Stakeholders and Roles','Customer, shop admin, developer/operator'),('5. Product Workflows','Customer, admin, and developer journeys'),('6. Functional Requirements','Current feature implementation'),('7. Architecture','Frontend, API, database, media, and deployment'),('8. Data and API Design','MongoDB entities, relationships, endpoint groups'),('9. Security and Privacy','Authentication, authorization, encryption, validation'),('10. Charges and KPIs','Daily rent estimate and business measures'),('11. UX and Reliability','Responsive design, performance practices, constraints'),('12. Deployment','Local setup, Vercel, Render, MongoDB, Cloudinary'),('13. Verification and Limits','Checks completed and remaining risks'),('14. Future Work','Potential enhancements'),('15. Conclusion','Project findings and summary'),('16. Technical References','Implementation and deployment sources')],[2.0,4.0],8.5)
callout('Implementation status','The report separates existing behavior from recommendations. It does not claim live production connectivity, payment processing, load testing, or external service verification.')
D.add_page_break()
H('1. Executive Summary')
para('RentEase is a responsive web platform for monthly furniture and appliance rentals. It connects customers with independent local shop owners, lets customers browse shop listings and submit rental requests, and gives each owner a private dashboard to manage inventory and rentals. A developer account provisions shop admins and views aggregate platform statistics.')
para('The project addresses the cost and logistics of purchasing household items for people who move for study or work. Its workflow includes a human confirmation step: customers select a shop and item, submit a delivery address, and the owner calls to approve or reject. The system starts a day-based rent estimate after approval. It does not collect payments or include resale.')
para('The application uses a React/Vite browser frontend, an Express/Node.js REST API, MongoDB through Mongoose, and Cloudinary for images. The intended production deployment separates the public frontend on Vercel from the API on Render. Each admin connects a Cloudinary account; its API secret is encrypted server-side, and the same account stores product and shop profile photos.')
table(['Element','Current implementation'],[('Users','Customers, shop admins, developer/operator'),('Business flow','Monthly rental requests; owner approval starts daily accrual'),('Frontend','React 18, Vite, JavaScript, responsive CSS, Lucide icons'),('Backend','Node.js, Express 4, REST endpoints, JWT sessions'),('Database','MongoDB/Mongoose'),('Images','Per-admin Cloudinary account for products and shop profile photos'),('Deployment target','Vercel frontend, Render API, MongoDB Atlas supported'),('Author','Azir Sam C; CSE; Web Development internship group')],[1.55,4.45],8.5)
# 2 Background
H('2. Background and Problem Statement');H('2.1 Context',2)
para('Students and working professionals often relocate for education or employment. Buying furniture and appliances can require a large initial outlay, transport planning, maintenance arrangements, and resale or disposal after another move. Local rental choices may be fragmented and difficult to compare.')
H('2.2 Problem addressed',2)
para('RentEase offers one storefront for discovering participating rental shops, comparing listed items and monthly prices, submitting delivery requests, and following active rentals. Each owner manages their own inventory and customers in a shop-scoped workspace. The developer/operator has a separate administrative overview.')
table(['Problem','Platform response'],[('High purchase cost','Monthly rental listings reduce initial commitment; no payment is collected in-app.'),('Relocation logistics','Customer supplies delivery address; selected shop coordinates by phone.'),('Flexible tenure','Active rentals can be returned operationally; formal tenure plans are not enforced in the app.'),('Local provider discovery','Shop cards expose complete storefronts and are ranked by distinct customer count.'),('Maintenance support','Owner can set maintenance status and add a tracking note; status is visible to customer.')],[1.65,4.35],8.3)
# 3 goals
H('3. Objectives and Scope');H('3.1 Primary objectives',2)
for x in ['Offer essential furniture and appliances through monthly rental listings.','Support a simple journey from local shop discovery to rental request.','Give every shop owner control over their products, orders, and rental statuses.','Calculate a day-based estimate only after an owner approves a request.','Provide a responsive experience on mobile and desktop browsers.']:bullet(x)
H('3.2 In scope',2)
for x in ['Browser application for customers, shop admins, and developer/operator.','Shop and product catalogs, product details, and customer rental requests.','Admin inventory management, request confirmation/rejection, maintenance, return, and delete actions.','Rental history, status, delivery address, and accrued rent estimate.','Cloudinary uploads for product and profile images using per-admin credentials.','MongoDB-backed REST API and Vercel/Render deployment configuration.']:bullet(x)
H('3.3 Out of scope',2)
for x in ['Native mobile applications, cross-border rentals, and AI-based pricing.','Online payments, payment instruments, invoices, or automatic billing.','Second-hand resale marketplace and automated logistics dispatch.','Live GPS tracking and connected-appliance telemetry.']:bullet(x)
callout('Scope boundary','RentEase tracks rental records and estimates accrued rent. Payment and physical delivery arrangements remain between the customer and shop.')
H('3.4 Assumptions and constraints',2)
para('The model assumes customers can access a modern browser and shop owners can call customers to coordinate. It depends on local inventory, reliable delivery and maintenance operations, and functioning owner Cloudinary accounts. Logistics costs, wear and tear, stock availability, and service coverage remain business constraints.')
# 4 roles
H('4. Stakeholders and Roles')
table(['Role','Responsibilities','Access boundary'],[('Customer','Browse shops and listings; submit requests; view own rentals; request return.','Public storefront plus own rental history.'),('Shop admin / owner','Complete profile; connect Cloudinary; manage listings; call customers; approve/reject; update maintenance and return status.','Own profile, products, and rentals scoped by owner ID.'),('Developer/operator','Provision shop admins and view aggregate platform counts and shop summaries.','Aggregate usage and admin provisioning; no general customer rental PII browser.'),('Deployment operator','Configure MongoDB, Render, Vercel, secrets, backups, and network rules.','Infrastructure access governed by provider account permissions.')],[1.2,2.7,2.1],8.1)
H('4.1 Authorization principle',2)
para('Protected API routes derive identity from a bearer token and check role middleware. Admin inventory and rental queries are scoped to the signed-in admin. Customer-only order routes prevent admins from placing rental requests as customers.')
# 5 flows
H('5. Product Workflows');H('5.1 Customer rental journey',2)
table(['Step','Action','System behavior'],[('1 Discover','Browse local shop cards.','Public API lists complete shop profiles, including profile image when present.'),('2 Select','Choose shop and browse items.','Only active and available items for that shop are shown.'),('3 Request','Sign in; enter phone and delivery address.','Creates pending rental under selected admin; item is reserved.'),('4 Confirm','Wait for owner phone call.','Owner approves or rejects; pending requests do not accrue rent.'),('5 Track','View status and accrued estimate.','Approval records start time; completed 24-hour periods accrue estimate.'),('6 Return','Coordinate collection with shop.','Customer may request return; owner marks returned to stop accrual.')],[.8,2.0,3.2],8)
H('5.2 Admin and developer journeys',2)
para('The developer creates an admin account. The owner signs in, completes the shop profile, connects Cloudinary, and adds products. New customer requests appear only in that owner’s dashboard. The owner calls the customer before approval. The developer can create or delete admins subject to open-rental safeguards and review aggregate customer/admin/shop statistics.')
# 6 features
H('6. Functional Requirements')
table(['ID','Requirement','Implemented behavior'],[('FR-01','Registration/login','Email and password; bcrypt hash; signed session token.'),('FR-02','Shop discovery','Complete shop profiles; ordering by distinct customers.'),('FR-03','Product browse','Furniture/Appliances categories, search, availability filters.'),('FR-04','Product details','Monthly rent, deposit, category, description, image.'),('FR-05','Rental request','Signed-in customer submits address and phone.'),('FR-06','Confirm/reject','Owner approves after phone call or rejects and deletes pending request.'),('FR-07','Rental lifecycle','Status, notes, maintenance, return, and history.'),('FR-08','Rent estimate','Monthly rent / 30 per complete 24-hour period after approval.'),('FR-09','Inventory','Create, edit, hide/list again, and delete with open-rental safeguards.'),('FR-10','Image upload','Product/profile uploads through authenticated per-admin Cloudinary.'),('FR-11','Developer console','Admin provisioning and aggregate platform/shop statistics.'),('FR-12','Password change','Signed-in flow requires current password; no email reset.'),('FR-13','Account deletion','Password confirmation and rental-state checks.'),('FR-14','Payments','Intentionally absent; no payment data collection.')],[.55,1.6,3.85],7.9)
# 7 architecture
H('7. Technical Architecture');H('7.1 Deployment layers',2)
table(['Layer','Technology','Responsibility'],[('Browser','React 18, Vite, JavaScript, CSS','Responsive UI, forms, navigation, rental summaries, API calls.'),('Frontend host','Vercel','Build and serve static assets from dist.'),('API','Node.js, Express 4, REST, JWT','Auth, business rules, validation, rate limits, upload orchestration.'),('Backend host','Render web service','Run API, store secrets, expose health check.'),('Database','MongoDB/Mongoose','Users, products, rental snapshots, encrypted credentials.'),('Media','Per-admin Cloudinary','Store each owner’s products and shop profile images.')],[1.05,1.9,3.05],8.2)
H('7.2 Request and deployment flow',2)
for x in ['React sends requests to the Render API using VITE_API_BASE_URL; this is public URL configuration only.','Render Express applies Helmet headers and checks the configured CORS origin allowlist.','JWT verification, role authorization, rate limits, and field validation run before protected business operations.','Mongoose queries use customer or owner scoping and persist to MongoDB.','For uploads, the API decrypts only the signed-in admin’s Cloudinary API secret and streams the validated image.','API responses return needed profile, product, or rental data; production errors remain generic.']:number(x)
para('For production, set VITE_API_BASE_URL in Vercel to the Render service URL without /api. Set CLIENT_ORIGIN in Render to the exact Vercel origin. Keep MONGODB_URI, JWT_SECRET, and CLOUDINARY_CREDENTIALS_KEY only in the backend environment. Do not rotate the encryption key without planning admin credential reconnection.')
# 8 data
H('8. Data and API Design');H('8.1 Core MongoDB entities',2)
table(['Entity','Selected fields','Purpose'],[('User','name, email, passwordHash, role, tokenVersion, phone, shop fields, storeLogo, encrypted Cloudinary connection','Customer/admin/developer identity and admin storefront. API secret ciphertext, IV, and tag are not returned to browser.'),('Product','owner, name, category, description, monthlyRent, deposit, image, imagePublicId, available, active','Admin-owned listing; owner scope controls management and media validation.'),('Rental','customer, owner, product, product snapshots, monthlyRent, deposit, status, notes, address, approvedAt, returnedAt','Order and rental lifecycle record; snapshots preserve checkout-era listing details.')],[.85,2.85,2.3],8)
H('8.2 Relationships and integrity',2)
for x in ['Each Product belongs to one admin through owner.','Each Rental references customer, owner, and product and copies name/category/rent/deposit snapshots.','Rental creation checks that the product is active, available, and owned by the selected shop.','Approval records approvedAt and reserves inventory; return records returnedAt and releases it.','Indexes support rental views by owner/status/date and product inventory by owner/active/availability/category.']:bullet(x)
H('8.3 API groups',2)
table(['Group','Representative routes','Access'],[('Auth','POST /api/auth/register, /login; GET /me; POST /change-password','Public auth or signed-in user.'),('Discovery','GET /api/admins; GET /api/products?adminId=...','Public restricted storefront data.'),('Customer rentals','POST /api/rentals; GET /mine; PATCH /:id/return','Customer only; own records.'),('Admin profile/images','GET/PATCH profile; POST/DELETE profile/image','Admin only.'),('Admin inventory/media','Cloudinary connect, POST /admin/images, CRUD /admin/products','Admin only; owner scope.'),('Admin rentals','GET /admin/rentals; PATCH/DELETE rentals; GET summary','Admin only; own records.'),('Developer','POST/DELETE developer/admins; GET developer/stats','Developer only.'),('Health','GET /api/health','Public readiness endpoint.')],[1.25,3.45,1.3],7.8)
# 9 Security
H('9. Security, Privacy, and Compliance');H('9.1 Technical safeguards',2)
table(['Control','Implementation'],[('Password storage','Passwords are hashed using bcrypt; plaintext passwords are not stored.'),('Sessions','JWT uses HS256 and expires in seven days. Token-version increments invalidate older sessions after password change.'),('Authorization','Role middleware protects customer, admin, and developer routes. Shop data queries are scoped by authenticated owner.'),('Cloudinary secrets','Per-admin API secrets are encrypted with AES-256-GCM and a server-only 32-byte key; the secret is decrypted only for the authenticated owner’s upload.'),('Rate limits','Global API, auth, account actions, Cloudinary connection, and image uploads have express-rate-limit controls. Default storage is per process.'),('Validation','JSON bodies are size-limited; values are normalized and bounded. Images are limited to 5 MB, restricted to JPEG/PNG/WebP/AVIF and signature-checked.'),('HTTP controls','Helmet security headers and CSP are enabled. CORS is configured with allowed origins; proxy hops are configurable for Render.'),('Errors','Production responses are generic; stack traces and raw provider/database errors are not returned.'),('Privacy','Public API is limited to storefront/product data; developer dashboard is aggregate-oriented.'),('Payments','No card, bank, or payment-provider data is collected.')],[1.35,4.65],8.1)
H('9.2 Operator responsibilities',2)
para('Application safeguards do not establish legal compliance by themselves. Before production, define the business identity, jurisdiction, privacy/support contacts, retention periods, account data deletion/export procedures, provider regions, backups, and incident response. Policy pages require those details and legal review.')
# 10
H('10. Rental Charges and Key Performance Indicators');H('10.1 Accrued rent estimate',2)
para('Pending requests accrue no rent. On approval, the API records approvedAt. For an active rental, the stop time is the current server time; for a returned rental, it is returnedAt. The day count is the number of complete 24-hour periods. The monthly rental snapshot is divided by 30.')
callout('Current formula','Accrued rent = rounded to 2 decimals [ floor((stop time - approvedAt) / 86,400,000) x (monthlyRent / 30) ]. This is an estimate, not payment collection, an invoice, a tax calculation, or a prorated contract.')
para('Maintenance and return-request statuses continue to accrue under current behavior. Charges stop when the owner marks the item returned. Reopening a returned rental sets a new approval time and begins a new accrual period.')
H('10.2 KPI definitions',2)
table(['KPI','Suggested definition','Current support'],[('Active rentals','Count of active, maintenance, or return-requested records.','Dashboard counts supported.'),('MRR','Sum monthly rent snapshots for active rentals.','Planning metric only; payments are not recorded.'),('Utilization','Unavailable/rented active listings divided by active listings.','Needs formalized reporting window and denominator.'),('Retention','Returning customers / eligible customers in a defined period.','Requires time-window/cohort analytics.'),('Maintenance resolution','Elapsed time from maintenance start to resolution.','No dedicated maintenance-start timestamp currently.')],[1.3,2.75,1.95],7.9)
para('Retention, MRR reconciliation, maintenance duration, and utilization trends are recommendations; the current app is not a complete analytics suite.')
# 11
H('11. User Experience, Performance, and Reliability');H('11.1 User experience',2)
para('The single-page application offers role-aware customer, admin, and developer workspaces, responsive layouts, mobile navigation, searchable product cards, profile images, confirmations for important actions, loading indicators, toast feedback, and a custom not-found page. Customer shop cards and the selected-shop banner show the owner’s profile image when available, with a visual fallback.')
H('11.2 Performance and reliability practices',2)
for x in ['Vite produces optimized production assets; catalog and admin record sets have query limits.','Cloudinary transformations bound image dimensions and request automatic format/quality.','Images use lazy loading where applicable; public shop data uses short caching while sensitive account data is not cached.','Rate limits reduce repeated traffic but do not guarantee a sub-three-second load time.','The default in-memory rate limiter is per process; multiple API instances require a shared store for consistent limits.']:bullet(x)
H('11.3 Constraints',2)
for x in ['Inventory accuracy depends on database state transitions; concurrency and high-volume booking should be load-tested.','Delivery, pickup, repair, damage inspection, and dispute handling remain manual shop operations.','Asset deletion is not coordinated in every account/listing removal flow; operators should plan Cloudinary housekeeping.','Losing CLOUDINARY_CREDENTIALS_KEY makes saved credentials unreadable and requires reconnecting admin accounts.','No live monitoring, alerting, production log aggregation, or disaster-recovery test is claimed.']:bullet(x)
# 12 deploy
H('12. Deployment and Operations');H('12.1 Local development',2)
for x in ['Install Node.js 20 or later and run npm install.','Copy .env.example to .env; configure MongoDB URI, JWT secret, developer credentials, and a 64-character hexadecimal Cloudinary encryption key.','Run npm run dev; Vite serves localhost:5173 and proxies /api to Express on port 10000.','Create an admin from the developer workspace, connect that admin’s Cloudinary account, and add products/profile image.']:number(x)
H('12.2 Render backend',2)
para('render.yaml defines a Node web service that runs npm ci --omit=dev and npm start, with /api/health as its health check. Configure MONGODB_URI, JWT_SECRET, DEVELOPER_EMAIL, DEVELOPER_PASSWORD, CLOUDINARY_CREDENTIALS_KEY, CLIENT_ORIGIN, NODE_ENV, and proxy settings in Render. Admin Cloudinary API secrets are per-admin and are not entered into Render.')
H('12.3 Vercel frontend',2)
para('Import the repository into Vercel as a Vite project. Build with npm run build and publish dist. Set VITE_API_BASE_URL to the Render base URL without /api. VITE-prefixed variables are included in client bundles and must never contain secrets.')
H('12.4 Production checklist',2)
for x in ['Restrict MongoDB Atlas network access to Render outbound IP ranges or an approved private route.','Set CLIENT_ORIGIN to the exact Vercel origin; explicitly list every allowed custom origin.','Confirm the Render health endpoint is healthy and review startup logs.','Verify customer discovery/request, admin approval, accrual display, return, and image upload with non-production accounts.','Back up MongoDB and CLOUDINARY_CREDENTIALS_KEY separately; document restore procedures.']:number(x)
# 13
H('13. Verification, Assumptions, and Limitations');H('13.1 Verification completed',2)
para('The production Vite build completed successfully in the workspace, and backend JavaScript syntax was checked during implementation. These checks confirm compilation and parsing, not live Vercel/Render connectivity, Cloudinary permissions, MongoDB availability, production response time, or end-to-end behavior against real accounts.')
H('13.2 Assumptions',2)
for x in ['Customers can use a modern browser and internet connection.','Shop admins can contact customers and arrange physical delivery and pickup.','Every admin maintains a valid Cloudinary product environment.','The deployment operator manages database access, HTTPS, secrets, backups, and provider accounts.']:bullet(x)
H('13.3 Limitations',2)
for x in ['Local inventory and logistics capacity may be limited.','No payment processing, delivery partner integration, or cross-border operations.','Tenure plans are not enforced as structured contract options.','No measured benchmark or formal penetration test is included.','Legal pages require operator-specific information and jurisdiction review.']:bullet(x)
# 14 future
H('14. Future Enhancements')
table(['Priority','Enhancement','Reason / prerequisite'],[('1','Automated tests for auth, owner isolation, rental lifecycle, upload, and edge cases.','Protect core behavior and access boundaries.'),('2','Shared rate-limit store, structured monitoring, alerts, and restore drills.','Support multiple instances and incident recovery.'),('3','Delivery/pickup scheduling and maintenance timestamps.','Measure operations and resolution KPIs.'),('4','Explicit tenure plans and rental extension requests.','Support clearer, flexible agreements.'),('5','Payment and auto-renewal integration if approved.','Requires separate financial, security, legal, and operational design.'),('6','Mobile applications, bundles, customization, or asset tracking.','Expand access and service options with privacy review.')],[.75,2.55,2.7],8)
H('15. Conclusion')
para('RentEase delivers the core workflow for a local multi-shop rental service: customers discover shops and request products, owners manage inventory and confirm requests by phone, and the application tracks status with a day-based rent estimate after approval. Developer tools support admin provisioning and aggregate platform oversight.')
para('The architecture targets a Vercel React frontend and Render Express API backed by MongoDB. Per-admin Cloudinary credentials support listing and owner-profile photos while remaining encrypted at rest. Before public launch, finish deployment configuration, end-to-end operational checks, monitoring and backup plans, and jurisdiction-specific policy review. Payment processing remains intentionally outside the current scope.')
H('16. Technical References')
para('Project files used to describe implementation and deployment:')
for x in ['README.md - product roles, workflows, environment variables, security, and deployment.','SETUP_GUIDE.md - local setup, admin onboarding, Cloudinary, and Vercel/Render steps.','package.json - project scripts and dependency declarations.','server/index.js - routes, authentication, authorization, validation, rental rules, uploads, and startup.','server/models.js - MongoDB schemas and indexes.','server/cloudinary.js - Cloudinary connection, image upload, and sanitized diagnostics.','src/main.jsx and src/styles.css - interfaces, user workflows, and responsive behavior.','render.yaml and .env.example - deployment service and environment template.']:bullet(x)
para('End of report',center=True)
# Keep table rows together
for t in D.tables:
 for row in t.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
D.save('RentEase_Project_Report.docx')
print(Path('RentEase_Project_Report.docx').resolve())
